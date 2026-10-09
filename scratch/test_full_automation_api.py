import requests
import openpyxl
import time
import os
from pathlib import Path

# 1. Create a sample Excel file with Gujarati headers matching user prompt
test_excel_path = Path("scratch/test_batch_10.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Villages"

headers = [
    "ક્રમ નંબર", "ઝોન", "જિલ્લાનું નામ", "જિલ્લાનો કોડ",
    "તાલુકાનું નામ", "તાલુકાનો કોડ", "ગામનું નામ", "ગામનો કોડ",
    "Electric (Kuva)", "Electric (BOR)", "Total Electric (BOR+Kuva)"
]
ws.append(headers)

sample_data = [
    [1, "Central", "Ahmedabad", "474", "Daskroi", "3801", "Bharkunda", "511639", 0, 42, 42],
    [2, "Central", "Ahmedabad", "474", "Sanand", "3802", "Changodar (Og)", "511624", 0, 15, 15],
    [3, "Central", "Ahmedabad", "474", "Sanand", "3802", "Hathipura", "511476", 0, 10, 10],
    [4, "Central", "Ahmedabad", "474", "Sanand", "3802", "Huka", "511635", 0, 8, 8],
    [5, "Central", "Ahmedabad", "474", "Dholka", "3803", "Kothadiya", "511910", 0, 12, 12],
    [6, "Central", "Ahmedabad", "474", "Detroj-Rampura", "3804", "Panar", "511494", 0, 25, 25],
    [7, "Central", "Ahmedabad", "474", "Dholka", "3803", "Parabdi", "511919", 0, 14, 14],
    [8, "Central", "Ahmedabad", "474", "Dholka", "3803", "Vejalka", "511750", 0, 18, 18],
    [9, "Central", "Anand", "475", "Khambhat", "3805", "Bhurakui", "516997", 0, 20, 20],
    [10, "Central", "Gandhinagar", "476", "Kalol", "3806", "(2) Rampura", "511130", 0, 0, 0]
]

for row in sample_data:
    ws.append(row)

wb.save(str(test_excel_path))
print(f"Created test Excel file: {test_excel_path} with {len(sample_data)} rows.")

# 2. Login to local Flask app
s = requests.Session()
login_res = s.post("http://127.0.0.1:5000/api/auth/login", json={
    "username": "admin",
    "password": "Admin@123"
})
print("Local App Login:", login_res.json())

# 3. Start automation with 5 workers
with open(test_excel_path, 'rb') as f:
    start_res = s.post("http://127.0.0.1:5000/api/automation/start", files={
        'file': ('test_batch_10.xlsx', f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    }, data={
        'concurrency': '5',
        'engine_mode': 'direct'
    })

print("Start API Response:", start_res.json())
job_id = start_res.json()['data']['job_id']

# 4. Poll status
t0 = time.time()
while True:
    time.sleep(0.5)
    st_res = s.get("http://127.0.0.1:5000/api/automation/status")
    st = st_res.json()['data']
    print(f"  [Status: {st['status']}] Processed: {st['processed_count']}/{st['total_villages']} | Success: {st['completed_count']} | Failed: {st['failed_count']} | Speed: {st['avg_speed_seconds']}s/vil | Throughput: {st['throughput_str']} | ETA: {st.get('estimated_remaining_time', '-')}")
    
    if st['status'] in ('completed', 'stopped', 'failed'):
        break

total_dur = round(time.time() - t0, 2)
print("=" * 60)
print(f"Automation finished in {total_dur} seconds!")
print(f"Total Villages: {st['total_villages']}")
print(f"Completed (Success): {st['completed_count']}")
print(f"Failed: {st['failed_count']}")
print(f"Avg Speed: {st['avg_speed_seconds']}s per village")
print(f"Throughput: {st['throughput_vpm']} villages/minute")
print("=" * 60)

# 5. Check Captured Results
print("\nCAPTURED RESULTS:")
for r in st['results']:
    print(f"  Row {r['row_num']}: {r['village_name']} ({r['village_code']}) -> {r['status']}: {r['generated_link'] or r['message']}")

# 6. Verify Output Excel download
dl_res = s.get(f"http://127.0.0.1:5000/api/automation/download/{job_id}")
print(f"\nDownload Output Excel HTTP status: {dl_res.status_code}, content-length: {len(dl_res.content)}")
out_test_path = Path(f"scratch/downloaded_{job_id}.xlsx")
with open(out_test_path, 'wb') as f:
    f.write(dl_res.content)

out_wb = openpyxl.load_workbook(str(out_test_path))
out_ws = out_wb.active
print(f"Output Excel sheet: '{out_ws.title}', max_row: {out_ws.max_row}, max_column: {out_ws.max_column}")
out_headers = [out_ws.cell(row=1, column=c).value for c in range(1, out_ws.max_column + 1)]
print(f"Output Headers: {out_headers}")

for r in range(2, out_ws.max_row + 1):
    v_n = out_ws.cell(row=r, column=7).value
    v_c = out_ws.cell(row=r, column=8).value
    lnk = out_ws.cell(row=r, column=out_ws.max_column - 2).value
    sta = out_ws.cell(row=r, column=out_ws.max_column - 1).value
    msg = out_ws.cell(row=r, column=out_ws.max_column).value
    print(f"  Row {r}: {v_n} ({v_c}) | Status: {sta} | Link: {lnk} | Msg: {msg}")
