import io
import os
from pathlib import Path
from flask import Blueprint, request, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from services.auth_service import login_required
from services.automation_service import automation_engine, build_output_excel
from config import Config

automation_bp = Blueprint('automation_bp', __name__)

@automation_bp.route('/api/automation/start', methods=['POST'])
@login_required
def start_automation():
    """Start an automation job from uploaded Excel file."""
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'No Excel file provided in upload.'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'success': False, 'error': 'No file selected.'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ('.xlsx', '.xls', '.csv'):
        return jsonify({'success': False, 'error': 'File must be an Excel spreadsheet (.xlsx, .xls) or CSV.'}), 400

    try:
        # Save uploaded file
        Config.ensure_dirs()
        upload_path = Config.AUTOMATION_UPLOAD_DIR / f"input_{int(os.times()[4]*100)}_{file.filename}"
        file.save(str(upload_path))

        # Optional parameters
        poll_interval = request.form.get('poll_interval', 0.25, type=float)
        max_poll_seconds = request.form.get('max_poll_seconds', 360, type=int)
        concurrency = request.form.get('concurrency', 5, type=int)
        engine_mode = request.form.get('engine_mode', 'direct').strip().lower()
        portal_url = request.form.get('portal_url', '').strip() or Config.PORTAL_URL or Config.MARKING_PORTAL_URL
        portal_username = request.form.get('portal_username', '').strip() or Config.PORTAL_USERNAME or Config.MARKING_PORTAL_USER
        portal_password = request.form.get('portal_password', '').strip() or Config.PORTAL_PASSWORD or Config.MARKING_PORTAL_PASSWORD
        headless_param = request.form.get('headless', 'true').strip().lower()
        headless = headless_param in ('true', '1', 'yes')
        auto_review_correct_val = request.form.get('auto_review_correct', 'true').strip().lower()
        auto_review_correct = auto_review_correct_val in ('true', '1', 'yes')
        force_fresh_val = request.form.get('force_fresh', 'false').strip().lower()
        force_fresh = force_fresh_val in ('true', '1', 'yes')
        resume_val = request.form.get('resume', 'false').strip().lower()
        resume = resume_val in ('true', '1', 'yes')

        options = {
            'poll_interval': poll_interval,
            'max_poll_seconds': max_poll_seconds,
            'concurrency': concurrency,
            'engine_mode': engine_mode,
            'portal_url': portal_url,
            'portal_username': portal_username,
            'portal_password': portal_password,
            'headless': headless,
            'auto_review_correct': auto_review_correct,
            'force_fresh': force_fresh,
            'resume': resume
        }

        job_info = automation_engine.start_job(
            file_path_or_storage=upload_path,
            filename=file.filename,
            options=options
        )

        return jsonify({
            'success': True,
            'data': job_info,
            'message': f"Automation started for {job_info['total_villages']} villages with {job_info.get('workers_count', concurrency)} concurrent workers."
        })

    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except RuntimeError as re:
        return jsonify({'success': False, 'error': str(re)}), 409
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@automation_bp.route('/api/automation/status', methods=['GET'])
@login_required
def get_automation_status():
    """Query live automation progress and results."""
    try:
        status_data = automation_engine.get_status()
        return jsonify({
            'success': True,
            'data': status_data
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@automation_bp.route('/api/automation/pause', methods=['POST'])
@login_required
def pause_automation():
    """Pause the running automation safely."""
    try:
        paused = automation_engine.pause_job()
        if paused:
            return jsonify({'success': True, 'message': 'Automation paused successfully.'})
        else:
            return jsonify({'success': False, 'error': 'Automation is not running or already paused.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@automation_bp.route('/api/automation/resume', methods=['POST'])
@login_required
def resume_automation():
    """Resume the paused automation safely."""
    try:
        resumed = automation_engine.resume_job()
        if resumed:
            return jsonify({'success': True, 'message': 'Automation resumed successfully.'})
        else:
            return jsonify({'success': False, 'error': 'Automation is not in paused state.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@automation_bp.route('/api/automation/stop', methods=['POST'])
@login_required
def stop_automation():
    """Stop the running automation."""
    try:
        stopped = automation_engine.stop_job()
        if stopped:
            return jsonify({'success': True, 'message': 'Automation stopping signal sent.'})
        else:
            return jsonify({'success': False, 'error': 'No running or paused automation job found.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@automation_bp.route('/api/automation/download/<job_id>', methods=['GET'])
@login_required
def download_automation_excel(job_id):
    """Download the newly generated output Excel containing original village data and Generated Village Page Links."""
    candidates = [
        f"Village_Generated_Links_{job_id}.xlsx",
        f"Village_KML_Links_{job_id}.xlsx",
        f"Village_KPI_Links_{job_id}.xlsx"
    ]
    filepath = None
    filename = candidates[0]

    for cand in candidates:
        cand_path = Config.AUTOMATION_UPLOAD_DIR / cand
        if cand_path.exists():
            filepath = cand_path
            filename = cand
            break

    if not filepath and automation_engine.output_filepath and automation_engine.output_filepath.exists():
        filepath = automation_engine.output_filepath
        filename = automation_engine.output_filename or filename

    if not filepath or not filepath.exists():
        return jsonify({'success': False, 'error': 'Output Excel file not found or still generating.'}), 404

    return send_file(
        str(filepath),
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@automation_bp.route('/api/automation/sample-template', methods=['GET'])
@login_required
def download_sample_template():
    """Download a ready-to-test sample Excel file containing 3 sample villages."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Villages"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Village Name", "District", "Priority Notes"]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_rows = [
        ["Amipura", "Ahmedabad", "Pilot well detection village"],
        ["Dhankuka", "Ahmedabad", "Subdistrict queue village"],
        ["Ashoknagar", "Ahmedabad", "GIS verification candidate"]
    ]

    for row in sample_rows:
        ws.append(row)

    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="sample_automation_villages.xlsx",
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
