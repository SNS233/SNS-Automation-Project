import os
import sys
import io
import time
import openpyxl
from pathlib import Path

# Ensure UTF-8 output on Windows
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from services.automation_service import (
    parse_village_excel,
    build_output_excel,
    find_excel_village_columns,
    VillageAutomationEngine
)
from services.browser_automation import PortalBrowserAutomator
from config import Config

def test_automation_workflow():
    print("==================================================")
    print("RUNNING VILLAGE AUTOMATION WORKFLOW TEST SUITE")
    print("==================================================")

    # 1. Test Gujarati Excel Detection & Parsing on actual uploaded Bhavnagar & Botad files
    print("\n[TEST 1] Testing Gujarati Column Detection on Real Excel Files...")
    for fname in ['uploads/excel/1791195066_Bhavnagar.xlsx', 'uploads/excel/1791216784_Botad.xlsx']:
        assert Path(fname).exists(), f"File {fname} must exist"
        wb = openpyxl.load_workbook(fname, data_only=True)
        ws = wb.active
        headers = [str(c.value) if c.value is not None else '' for c in ws[1]]
        mapping = find_excel_village_columns(headers)
        
        # Verify Village Name (ગામનું નામ) and Village Code (ગામનો કોડ)
        v_name_header = headers[mapping['village_name_idx']]
        v_code_header = headers[mapping['village_code_idx']]
        print(f"  File: {Path(fname).name}")
        print(f"    Detected Village Name col ({mapping['village_name_idx']}): '{v_name_header}'")
        print(f"    Detected Village Code col ({mapping['village_code_idx']}): '{v_code_header}'")
        
        assert "ગામ" in v_name_header and "નામ" in v_name_header, f"Expected village name header, got {v_name_header}"
        assert "ગામ" in v_code_header and "કોડ" in v_code_header, f"Expected village code header, got {v_code_header}"
        # Ensure district code was not confused
        assert "જિલ્લા" not in v_code_header and "તાલુકા" not in v_code_header

        rows = parse_village_excel(fname)
        print(f"    Parsed {len(rows)} villages successfully. Sample: {rows[0]['village_name']} (Code: {rows[0]['village_code']})")
        assert len(rows) > 0, "Should have parsed at least 1 row"

    # 2. Test Output Excel Builder: Retaining ALL original columns + Generated Link, Generation Status, Error/Message
    print("\n[TEST 2] Testing Output Excel Builder (Preserving Original Columns + Hyperlinks + Blank Failures)...")
    test_input_path = Path("uploads/excel/1791216784_Botad.xlsx")
    test_out_path = Config.AUTOMATION_UPLOAD_DIR / "test_botad_output.xlsx"

    sample_results = [
        {
            "row_num": 2,
            "village_name": "Pipariya",
            "village_code": "511852",
            "generated_link": "https://generate.snsies.com/villages/30/review",
            "status": "SUCCESS",
            "message": "Generated successfully"
        },
        {
            "row_num": 3,
            "village_name": "Nabhoi",
            "village_code": "511853",
            "generated_link": "",
            "status": "FAILED",
            "error": "Timeout waiting for generation",
            "message": "Timeout waiting for generation"
        },
        {
            "row_num": 4,
            "village_name": "Dhadhoda",
            "village_code": "511854",
            "generated_link": "https://generate.snsies.com/villages/32/review",
            "status": "SKIPPED",
            "message": "Already completed"
        }
    ]

    build_output_excel(test_out_path, sample_results, input_filepath=test_input_path)
    assert test_out_path.exists(), "Output Excel file must be created"

    out_wb = openpyxl.load_workbook(str(test_out_path))
    out_ws = out_wb.active
    out_headers = [str(cell.value) for cell in out_ws[1]]
    print(f"  Output Excel Headers ({len(out_headers)} columns):")
    print(f"    {out_headers}")

    # Verify original columns exist
    assert "ક્રમ નંબર" in out_headers
    assert "ગામનું નામ" in out_headers
    assert "ગામનો કોડ" in out_headers

    # Verify the 3 appended columns
    assert "Generated Link" in out_headers
    assert "Generation Status" in out_headers
    assert "Error/Message" in out_headers

    link_col_idx = out_headers.index("Generated Link") + 1
    status_col_idx = out_headers.index("Generation Status") + 1
    msg_col_idx = out_headers.index("Error/Message") + 1

    # Verify Row 2 (SUCCESS: Link present with hyperlink)
    row2_link = out_ws.cell(row=2, column=link_col_idx).value
    row2_hyperlink = out_ws.cell(row=2, column=link_col_idx).hyperlink
    row2_status = out_ws.cell(row=2, column=status_col_idx).value
    print(f"  Row 2 Data (Success): Link='{row2_link}', Status='{row2_status}'")
    assert row2_link == "https://generate.snsies.com/villages/30/review"
    assert row2_status == "SUCCESS"
    assert row2_hyperlink is not None and row2_hyperlink.target == "https://generate.snsies.com/villages/30/review"

    # Verify Row 3 (FAILED: Link must be BLANK, status FAILED)
    row3_link = out_ws.cell(row=3, column=link_col_idx).value
    row3_status = out_ws.cell(row=3, column=status_col_idx).value
    row3_msg = out_ws.cell(row=3, column=msg_col_idx).value
    print(f"  Row 3 Data (Failed): Link='{row3_link}', Status='{row3_status}', Msg='{row3_msg}'")
    assert not row3_link or row3_link == "", "Failed row MUST have blank Generated Link"
    assert row3_status == "FAILED"
    assert "Timeout" in row3_msg

    # Verify Row 4 (SKIPPED)
    row4_status = out_ws.cell(row=4, column=status_col_idx).value
    assert row4_status == "SKIPPED"

    # Cleanup test output file
    if test_out_path.exists():
        test_out_path.unlink()

    # 3. Test Real Browser Automator Headless Driver Launch
    print("\n[TEST 3] Testing Selenium Chrome Driver Initializer (Headless)...")
    automator = PortalBrowserAutomator(portal_url="http://127.0.0.1:5000", headless=True)
    driver = automator.start()
    assert driver is not None, "Chrome driver instance should be initialized"
    print(f"  Chrome driver launched successfully: {driver.session_id}")
    automator.close()
    assert automator.driver is None, "Driver should be quit and closed"
    print("  Chrome driver closed cleanly.")

    # 4. Test VillageAutomationEngine State and Metrics Snapshot
    print("\n[TEST 4] Testing Village Automation Engine KPIs and State Snapshot...")
    engine = VillageAutomationEngine()
    st = engine.get_status()
    print(f"  Engine Status snapshot: Total={st['total_villages']}, Processed={st['processed_count']}, Status={st['status']}")
    assert 'total_villages' in st
    assert 'completed_count' in st
    assert 'failed_count' in st
    assert 'remaining_count' in st
    assert 'processed_count' in st
    assert 'progress_percentage' in st
    assert 'workers_count' in st
    assert 'avg_speed_seconds' in st
    assert 'throughput_str' in st

    # 5. Test Engine Pause and Resume Lifecycle
    print("\n[TEST 5] Testing Village Automation Engine Pause / Resume Lifecycle...")
    # Simulate a running state to test pause/resume events
    engine.status = "running"
    engine.pause_event.set()
    engine.active_villages = {1: "TestVillage"}

    # Test Pause
    paused = engine.pause_job()
    assert paused is True, "pause_job() should return True when running"
    assert engine.status == "paused"
    assert engine.is_paused is True
    assert not engine.pause_event.is_set(), "pause_event should be cleared on pause"
    st_paused = engine.get_status()
    assert st_paused['status'] == "paused"
    print(f"  Pause verified: Status='{st_paused['status']}', PauseEventSet={engine.pause_event.is_set()}")

    # Test Resume
    resumed = engine.resume_job()
    assert resumed is True, "resume_job() should return True when paused"
    assert engine.status == "running"
    assert engine.is_paused is False
    assert engine.pause_event.is_set(), "pause_event should be set on resume"
    st_resumed = engine.get_status()
    assert st_resumed['status'] == "running"
    print(f"  Resume verified: Status='{st_resumed['status']}', PauseEventSet={engine.pause_event.is_set()}")

    # Test Stop
    stopped = engine.stop_job()
    assert stopped is True
    assert engine.status == "stopped"
    engine.active_villages.clear()
    engine.status = "idle"
    print(f"  Stop verified: Status='{engine.status}'")

    print("\n==================================================")
    print("ALL TESTS PASSED WITH 100% COMPLIANCE!")
    print("==================================================")

if __name__ == "__main__":
    test_automation_workflow()
