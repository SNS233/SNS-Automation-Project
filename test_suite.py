import io
import os
import openpyxl
import pandas as pd
from app import create_app
import db
from services import kml_service, excel_service, village_service

def run_tests():
    print("==================================================")
    print("RUNNING MASTER PORTAL VERIFICATION TEST SUITE")
    print("==================================================")

    # 1. DB Health Check
    health = db.check_db_health()
    print("[TEST 1] SQL Server Health Check:", health)
    assert health['connected'] is True, "SQL Server should be connected"

    # 2. Test Sample Excel Generation
    print("\n[TEST 2] Generating Test Excel File with Sample Villages...")
    sample_wb = openpyxl.Workbook()
    ws = sample_wb.active
    ws.title = "Villages"

    headers = [
        "Zone", "District_Name", "Subdistrict_Name", "Village name", "Village Code",
        "Geb Connection", "Google search original Nos", "Google edited Nos and final KML",
        "Total", "Dugwells", "Bores", "KML link", "KML status (in progress/done)", "Remark"
    ]
    ws.append(headers)

    sample_villages = [
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Sarsav", "504543", "GEB-101", 12, 12, 12, 4, 8, "", "Pending", "Priority village"],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Saroli", "504544", "GEB-102", 5, 5, 5, 2, 3, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Parvath", "504545", "GEB-103", 8, 8, 8, 1, 7, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Navabhaga", "504546", "GEB-104", 15, 15, 15, 5, 10, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Mondhri", "504547", "GEB-105", 3, 3, 3, 1, 2, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Masota", "504548", "GEB-106", 9, 9, 9, 3, 6, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Dadhav", "504549", "GEB-107", 7, 7, 7, 2, 5, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "Biladiya", "504550", "GEB-108", 11, 11, 11, 4, 7, "", "Pending", ""],
        ["GPIC", "Sabar Kantha", "Vijaynagar", "InvalidVillageNoCode", "", "GEB-109", 0, 0, 0, 0, 0, "", "Pending", "Missing code test"]
    ]

    for row in sample_villages:
        ws.append(row)

    sample_excel_path = "test_sample_villages.xlsx"
    sample_wb.save(sample_excel_path)
    print(f"Created '{sample_excel_path}' with {len(sample_villages)} test rows.")

    # 3. Test Preview
    print("\n[TEST 3] Testing File Preview Service...")
    preview = excel_service.preview_uploaded_file(sample_excel_path)
    print("Preview headers detected:", preview['headers'])
    print("Suggested mappings:", preview['suggested_mappings'])
    print("Missing required fields:", preview['missing_required'])
    assert len(preview['missing_required']) == 0, "All required canonical fields should be auto-mapped"

    # 4. Test Full Batch Import
    print("\n[TEST 4] Testing Batch Import into SQL Server...")
    # Clean previous test entries if any
    db.execute_non_query("DELETE FROM Villages WHERE district_name = 'Sabar Kantha'")
    
    import_result = excel_service.parse_and_import_file(
        file_storage_or_path=sample_excel_path,
        column_mappings=preview['suggested_mappings'],
        duplicate_mode='skip',
        batch_filename='test_sample_villages.xlsx'
    )
    print("Import Result:", import_result)
    assert import_result['inserted_count'] == 8, f"Expected 8 valid inserted villages, got {import_result['inserted_count']}"
    assert import_result['invalid_rows'] == 1, f"Expected 1 invalid row (missing code), got {import_result['invalid_rows']}"
    assert import_result['error_report_filename'] is not None, "Error report should be generated"

    # 5. Test Querying Villages
    print("\n[TEST 5] Querying Imported Villages...")
    paged = village_service.get_villages_paged(page=1, page_size=10, search='Sarsav')
    assert paged['total_count'] >= 1, "Sarsav should be found"
    sarsav = paged['items'][0]
    print(f"Found Sarsav village ID: {sarsav['id']}, Code: {sarsav['village_code']}, Dugwells: {sarsav['dugwells']}, Bores: {sarsav['bores']}")

    # 6. Test KML Parsing & Linking
    print("\n[TEST 6] Testing KML Parser & Storage...")
    sample_kml_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <kml xmlns="http://www.opengis.net/kml/2.2">
      <Document>
        <name>Sarsav Village Wells</name>
        <Placemark>
          <name>Tubewell 1</name>
          <description>Borewell 150ft deep</description>
          <Point><coordinates>73.01234,23.45678,0</coordinates></Point>
        </Placemark>
        <Placemark>
          <name>Dugwell 1</name>
          <description>Open village community well</description>
          <Point><coordinates>73.01500,23.45800,0</coordinates></Point>
        </Placemark>
        <Placemark>
          <name>Tubewell 2</name>
          <description>Farm borewell</description>
          <Point><coordinates>73.01800,23.46000,0</coordinates></Point>
        </Placemark>
        <Placemark>
          <name>Agricultural Field Zone</name>
          <Polygon>
            <outerBoundaryIs>
              <LinearRing>
                <coordinates>
                  73.010,23.450,0 73.020,23.450,0 73.020,23.460,0 73.010,23.460,0 73.010,23.450,0
                </coordinates>
              </LinearRing>
            </outerBoundaryIs>
          </Polygon>
        </Placemark>
      </Document>
    </kml>
    """

    kml_res = kml_service.save_kml_for_village(
        village_id=sarsav['id'],
        file_storage_or_bytes=sample_kml_xml.encode('utf-8'),
        original_filename='Sarsav_504543_wells.kml',
        village_code=sarsav['village_code']
    )
    print("KML Saved Result:", kml_res)
    assert kml_res['point_count'] == 3, f"Expected 3 points, got {kml_res['point_count']}"
    assert kml_res['polygon_count'] == 1, f"Expected 1 polygon, got {kml_res['polygon_count']}"
    assert kml_res['detected_borewells'] == 2, f"Expected 2 detected borewells, got {kml_res['detected_borewells']}"
    assert kml_res['detected_dugwells'] == 1, f"Expected 1 detected dugwell, got {kml_res['detected_dugwells']}"

    # 7. Test Village Status & Queue Workflow
    print("\n[TEST 7] Testing Queue Progression & Status Update...")
    village_service.update_village(
        village_id=sarsav['id'],
        data={
            'bores': 3,
            'dugwells': 1,
            'total_wells': 4,
            'kml_status': 'Done',
            'remark': 'Verified against KML and GIS satellite layer.'
        }
    )
    updated_sarsav = village_service.get_village_by_id(sarsav['id'])
    assert updated_sarsav['kml_status'] == 'Done'
    assert updated_sarsav['total_wells'] == 4
    print("Updated Sarsav status:", updated_sarsav['kml_status'], "History entries:", len(updated_sarsav['history']))

    # 8. Test Queue Navigation
    print("\n[TEST 8] Testing Queue Navigation...")
    queue_nav = village_service.get_queue_navigation(current_id=sarsav['id'])
    print(f"Current: {queue_nav['current_village']['village_name']}, Next Pending ID: {queue_nav['next_village_id']}, Pending Count: {queue_nav['pending_count']}")
    assert queue_nav['next_village_id'] is not None, "Next pending village should be present"

    # 9. Test Excel Export
    print("\n[TEST 9] Testing Excel Export Generation...")
    export_stream = excel_service.generate_excel_export(base_url="http://localhost:5000")
    exported_bytes = export_stream.getvalue()
    assert len(exported_bytes) > 1000, "Exported Excel should contain data"

    # Inspect exported workbook structure
    exp_wb = openpyxl.load_workbook(io.BytesIO(exported_bytes))
    print("Exported Sheets:", exp_wb.sheetnames)
    assert "Village Data" in exp_wb.sheetnames
    assert "Summary Report" in exp_wb.sheetnames
    data_sheet = exp_wb["Village Data"]
    exp_headers = [cell.value for cell in data_sheet[1]]
    assert len(exp_headers) >= 14, "Exported headers must contain all columns!"

    # Check that Sarsav has the working KML link
    sarsav_row = None
    for row in data_sheet.iter_rows(min_row=2, values_only=True):
        if "504543" in [str(c) for c in row]:  # Village Code
            sarsav_row = row
            break
    assert sarsav_row is not None, "Sarsav row should exist in export"
    kml_links = [str(c) for c in sarsav_row if "/api/kml/download/" in str(c)]
    assert len(kml_links) > 0, "KML Link should contain working download URL"


    # 10. Dashboard Metrics Verification
    print("\n[TEST 10] Testing Dashboard Dynamic Metrics...")
    metrics = village_service.get_dashboard_metrics()
    print("Summary Metrics:", metrics['summary'])
    assert metrics['summary']['total_villages'] >= 8
    assert metrics['summary']['completed_villages'] >= 1
    assert metrics['summary']['kml_files_count'] >= 1

    print("\n==================================================")
    print("ALL TESTS PASSED SUCCESSFULLY!")
    print("==================================================")

if __name__ == "__main__":
    run_tests()
